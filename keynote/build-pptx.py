"""
Build an editable PowerPoint version of the keynote.

Keynote opens .pptx natively and keeps every text box and image editable,
which a PDF import cannot do. Layout mirrors deck.html; the canvas is
1280x720 px, and 1 px == 1/96 in so px * 9525 == EMU.
"""
import os
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE

PX = 9525                      # EMU per CSS pixel at 96 dpi
def E(px): return Emu(int(round(px * PX)))

PAPER  = RGBColor(0xF6, 0xF5, 0xF2)
INK    = RGBColor(0x17, 0x19, 0x1D)
MUTED  = RGBColor(0x5C, 0x61, 0x69)
ACCENT = RGBColor(0x3F, 0x61, 0x89)
LINE   = RGBColor(0xDF, 0xDC, 0xD5)
FRAME  = RGBColor(0xEA, 0xE7, 0xE0)
EDGE   = RGBColor(0xD8, 0xD4, 0xCB)
FTEXT  = RGBColor(0x9A, 0x95, 0x8A)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
WHITEISH = RGBColor(0xB8, 0xBC, 0xC2)

FONT = "Helvetica Neue"        # present on every Mac; closest match to Inter

HERE   = os.path.dirname(os.path.abspath(__file__))
PHOTOS = os.path.join(HERE, "photos")
CROPS  = os.path.join(HERE, ".crops")
os.makedirs(CROPS, exist_ok=True)


def crop_to(name, w, h, scrim=False):
    """Centre-crop a photo to the frame's aspect ratio. Returns a path or None."""
    src = os.path.join(PHOTOS, name)
    if not os.path.exists(src):
        return None
    im = Image.open(src).convert("RGB")
    tw, th = w * 2, h * 2                       # 2x for a crisp result
    scale = max(tw / im.width, th / im.height)
    im = im.resize((max(tw, int(im.width * scale)), max(th, int(im.height * scale))),
                   Image.LANCZOS)
    left, top = (im.width - tw) // 2, (im.height - th) // 2
    im = im.crop((left, top, left + tw, top + th))

    if scrim:
        # Same two gradients deck.html lays over the slide 3 hero, baked in
        # so the title stays legible without a transparency effect in pptx.
        ov = Image.new("RGBA", im.size, (0, 0, 0, 0))
        d = ImageDraw.Draw(ov)
        for x in range(im.width):
            t = x / im.width
            a = 0.52 - 0.32 * min(t / 0.46, 1.0) if t < 0.46 else \
                max(0.0, 0.20 * (1 - (t - 0.46) / 0.26))
            d.line([(x, 0), (x, im.height)], fill=(12, 14, 18, int(a * 255)))
        im = Image.alpha_composite(im.convert("RGBA"), ov)
        ov = Image.new("RGBA", im.size, (0, 0, 0, 0))
        d = ImageDraw.Draw(ov)
        for y in range(im.height):
            t = y / im.height
            a = 0.04 + 0.06 * (t / 0.40) if t < 0.40 else 0.10 + 0.48 * ((t - 0.40) / 0.60)
            d.line([(0, y), (im.width, y)], fill=(12, 14, 18, int(a * 255)))
        im = Image.alpha_composite(im, ov).convert("RGB")

    out = os.path.join(CROPS, f"{os.path.splitext(name)[0]}_{w}x{h}.jpg")
    im.save(out, "JPEG", quality=92)
    return out


def track(run, em, size_pt):
    """Letter-spacing, which python-pptx has no API for."""
    run.font._rPr.set("spc", str(int(round(em * size_pt * 100))))


def est_height(paras, w, size, leading, space_before):
    """Rough wrapped height in px - only needs to be generous, not exact."""
    cpl = max(1, int(w / (size * 0.50)))
    lines = sum(max(1, -(-len(t) // cpl)) for t in paras)
    return lines * size * leading + (len(paras) - 1) * space_before + size


def text(slide, x, y, w, s, size, color, weight=400, spacing=0.0,
         leading=1.2, align=PP_ALIGN.LEFT, space_before=0):
    paras = s if isinstance(s, list) else [s]
    h = est_height(paras, w, size, leading, space_before)
    box = slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, para_text in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = leading
        if i > 0 and space_before:
            p.space_before = Pt(space_before * 0.75)
        r = p.add_run()
        r.text = para_text
        r.font.size = Pt(size * 0.75)          # px -> pt
        r.font.name = FONT
        r.font.bold = weight >= 600
        r.font.color.rgb = color
        if spacing:
            track(r, spacing, size * 0.75)
    return box


def photo(slide, x, y, w, h, name, label, scrim=False):
    path = crop_to(name, w, h, scrim)
    if path:
        slide.shapes.add_picture(path, E(x), E(y), E(w), E(h))
        return
    box = slide.shapes.add_shape(1, E(x), E(y), E(w), E(h))   # rectangle
    box.fill.solid(); box.fill.fore_color.rgb = FRAME
    box.line.color.rgb = EDGE; box.line.width = Pt(0.75)
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label
    r.font.size = Pt(10.5 * 0.75); r.font.name = FONT; r.font.color.rgb = FTEXT
    track(r, 0.12, 10.5 * 0.75)


def rule(slide, x, y, w, h, color):
    s = slide.shapes.add_shape(1, E(x), E(y), E(w), E(h))
    s.fill.solid(); s.fill.fore_color.rgb = color
    s.line.fill.background(); s.shadow.inherit = False


prs = Presentation()
prs.slide_width, prs.slide_height = E(1280), E(720)
BLANK = prs.slide_layouts[6]

def new_slide():
    s = prs.slides.add_slide(BLANK)
    fill = s.background.fill
    fill.solid(); fill.fore_color.rgb = PAPER
    return s

EYEBROW = dict(size=10.5, weight=500, spacing=0.16)
TITLE   = dict(size=46, weight=600, spacing=-0.025, leading=1.08)
BODY    = dict(size=16.5, weight=400, leading=1.78, space_before=15)
NUM     = dict(size=10.5, weight=500, spacing=0.14)
NUMR    = dict(NUM, align=PP_ALIGN.RIGHT)

# ---------------------------------------------------------------- 01 cover
s = new_slide()
photo(s, 668, 0, 612, 720, "kevin.jpg", "[ ADD KEVIN PHOTO HERE ]")
text(s, 88, 214, 560, "ARCHBISHOP MCCARTHY HIGH SCHOOL   ·   10 YEAR DREAM KEYNOTE",
     color=MUTED, **EYEBROW)
text(s, 88, 246, 500, ["Kevin", "Hernandez"], size=62, weight=600,
     spacing=-0.03, leading=1.02, color=INK)
text(s, 88, 404, 500, "Build a life that feels like mine.",
     size=25, weight=500, spacing=-0.012, leading=1.35, color=ACCENT)
rule(s, 88, 472, 34, 2, ACCENT)
text(s, 88, 500, 404, "In 10 years I hope to be back in Miami, running my own "
     "business, with a family and a life I’m proud of.", color=MUTED, **BODY)
text(s, 88, 646, 200, "01 / 05", color=RGBColor(0xA6,0xA2,0x99), **NUM)

# ------------------------------------------------------------ 02 education
s = new_slide()
photo(s, 620, 208, 572, 312, "campus.jpg", "COLLEGE CAMPUS")
text(s, 88, 196, 470, "EDUCATION", color=MUTED, **EYEBROW)
text(s, 88, 228, 470, "The Next Step", color=INK, **TITLE)
text(s, 88, 304, 470, [
    "After McCarthy I want to go to college and study business. I’m most "
    "interested in the entrepreneurship side of it — finance, marketing, and "
    "how a company actually runs day to day.",
    "I also want to use college to meet people and get real experience, not "
    "just sit in class. An internship, a job, whatever I can get. And if I "
    "come up with an idea I believe in, I’d like to try starting something "
    "small while I’m still in school."], color=MUTED, **BODY)
text(s, 88, 646, 200, "02 / 05", color=RGBColor(0xA6,0xA2,0x99), **NUM)

# --------------------------------------------------------- 03 career/miami
s = new_slide()
photo(s, 0, 0, 1280, 392, "miami.jpg", "MIAMI SKYLINE", scrim=True)
have_miami = os.path.exists(os.path.join(PHOTOS, "miami.jpg"))
text(s, 88, 262, 760, "CAREER   ·   MIAMI, FLORIDA",
     color=WHITEISH if have_miami else MUTED, **EYEBROW)
text(s, 88, 290, 760, "Building My Own Thing",
     color=WHITE if have_miami else INK, **TITLE)
text(s, 88, 452, 520, [
    "Ten years out, I picture myself in Miami running my own business. I don’t "
    "know exactly what it will be yet — that’s the honest answer. I just know "
    "I want it to be mine.",
    "I’m fine with working a lot, especially early on. What I really want is "
    "control over my own time, enough money that I’m not worried about money, "
    "and something I built that I’d put my name on."], color=MUTED, **BODY)
text(s, 992, 676, 200, "03 / 05", color=RGBColor(0xA6,0xA2,0x99), **NUMR)

# ----------------------------------------------------------- 04 life/hobby
s = new_slide()
photo(s, 88, 88, 540, 544, "family.jpg", "FAMILY")
text(s, 704, 196, 488, "LIFE OUTSIDE OF WORK", color=MUTED, **EYEBROW)
text(s, 704, 228, 488, "The Life I Want", color=INK, **TITLE)
text(s, 704, 304, 488, [
    "Outside of work I keep picturing something pretty normal, honestly. A "
    "wife, two kids, a dog, and a house with enough room for people to come "
    "over. Sundays with family. The same friends I have now, still around.",
    "I want to stay in the gym, get to the beach when I can, keep up with my "
    "teams, and probably spend too much on a car at some point."],
    color=MUTED, **BODY)
text(s, 992, 646, 200, "04 / 05", color=RGBColor(0xA6,0xA2,0x99), **NUMR)

# -------------------------------------------------------- 05 travel/reunion
s = new_slide()
photo(s, 0, 0, 612, 720, "barcelona.jpg", "BARCELONA, SPAIN")
text(s, 696, 150, 496, "THE REUNION", color=MUTED, **EYEBROW)
text(s, 696, 182, 496, "See You in 10", color=INK, **TITLE)
text(s, 696, 258, 470, [
    "If it goes the way I hope: I’m living in Miami, running my own business, "
    "and I pull up with my wife, our two kids, and probably the dog too.",
    "By then I want to have made it to Barcelona — Spain is the one trip I "
    "really want to take. Comfortable, not rich. Proud of what I built, and "
    "still close with the same people I care about now."], color=MUTED, **BODY)
rule(s, 696, 512, 462, 1, LINE)
text(s, 696, 540, 462, "I know it won’t go exactly like this. I just want to "
     "look back and know it was all mine.",
     size=20, weight=500, spacing=-0.014, leading=1.5, color=ACCENT)
text(s, 992, 646, 200, "05 / 05", color=RGBColor(0xA6,0xA2,0x99), **NUMR)

out = os.path.join(HERE, "Kevin-Hernandez-10-Year-Dream-Keynote.pptx")
prs.save(out)
print("wrote", out, f"({round(os.path.getsize(out)/1024)} KB, {len(prs.slides.__iter__.__self__._sldIdLst)} slides)")
